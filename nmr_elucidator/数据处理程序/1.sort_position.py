# -*- coding: utf-8 -*-
"""仅对第一个 sheet 的行排序：不带′的单独数字按数字升序（OCH3 特殊值排在该组最后），
带′的先按′数量升序、数量相同再按最后一个数字升序；
排序后删除最后一个数字之前的所有字符串（OCH3 特殊值保留原样），′ 统一为半角 '。
其它 sheet 内容原样不动，输出到 原文件名_已排序.xlsx。"""
import os
import pandas as pd
from openpyxl import load_workbook

SRC = r'Y:\LXH数据处理\新化合物处理\程序解析\clxh122解析\clxh122_归属表.xlsx'

OUT = os.path.splitext(SRC)[0] + '_已排序.xlsx'

wb = load_workbook(SRC)
ws = wb.worksheets[0]                       # 只处理第一个 sheet
header = [c.value for c in ws[1]]
df = pd.DataFrame(list(ws.values)[1:], columns=header)

pos_col = next((c for c in df.columns
                if str(c).strip().lower() in ('position', 'pos.', 'pos', 'no.', 'no')), None)
if pos_col is None:
    raise KeyError(f"未找到位置列，现有列名: {df.columns.tolist()}")

# 表格尾部有大量全空行（所有列均为空），先剔除再排序
df = df.dropna(how='all').reset_index(drop=True)

# 统一转为半角撇号（覆盖全角′、全角'等变体）
s = (df[pos_col].astype(str).str.strip()
       .str.replace('\u2032', "'", regex=False)
       .str.replace('\uff07', "'", regex=False))

# 提取末尾数字和′数量（全程向量化，无 for 循环）
is_special = s.str.contains('OCH3', case=False)
clean = s.str.extract(r"(\d+'*)$")[0]          # 删除最后一个数字之前的所有字符串
s = s.where(is_special, clean)                 # OCH3 特殊值保留原样
df[pos_col] = s

lastnum = s.str.extract(r"(\d+)'*$")[0].astype(int)
nprime = s.str.count("'")

# 不带′的排最前（组内特殊值排最后、普通数字按数字升序），其后按′数量、数字升序
order = (df.assign(hasprime=(nprime > 0).astype(int), nprime=nprime,
                   is_special=is_special.astype(int), lastnum=lastnum)
           .sort_values(['hasprime', 'nprime', 'is_special', 'lastnum']).index)

# 按排序结果原位回写第一个 sheet（其它 sheet 不受影响）
for new_i, old_i in enumerate(order, start=2):
    for j, v in enumerate(df.loc[old_i], start=1):
        ws.cell(row=new_i, column=j, value=v)

wb.save(OUT)
print(df.loc[order, pos_col].to_string())

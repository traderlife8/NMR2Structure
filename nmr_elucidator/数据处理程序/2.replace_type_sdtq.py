# -*- coding: utf-8 -*-
import time
import numpy as np
import pandas as pd
from openpyxl import load_workbook

t0 = time.perf_counter()

# 不动原始"类型"列，新建 type 列存放 13C NMR 多重性小写字母（s/d/t/q），
# ppm 列保留一位小数，类型列为空的行批量删除。OCH3 行通过位置列中的 "OCH3" 标记识别。
# 用 openpyxl 在第一个 Sheet 上就地修改，原字体/列宽等格式原样保留；其余 Sheet 不动。
SRC = r'x:\OCF投稿\2026-09-26数据处理\CP7.xlsx'

wb = load_workbook(SRC)
ws = wb.worksheets[0]

header = {str(c.value).strip().lower(): c.column for c in ws[1] if c.value is not None}

def find_col(candidates):
    for n in candidates:
        col = header.get(n.lower())
        if col is not None:
            return col
    raise KeyError(f"未找到列，候选: {candidates}，现有表头: {list(header)}")

pos_col = find_col(('position', 'pos.', 'pos', 'no.', 'no'))
type_col = find_col(('类型', 'types'))
ppm_col = find_col(('ppm', '13C NMR', 'δc'))
type_new_col = header.get('type', (max(header.values()) + 1))
ws.cell(row=1, column=type_new_col, value='type')

# 数据区读入 DataFrame，df 第 i 行对应工作表第 i+2 行（保持对齐，便于向量化计算）
data = pd.DataFrame(ws.iter_rows(min_row=2, values_only=True))

m = {'C': 's', 'CH': 'd', 'CH2': 't', 'CH3': 'q', 'OCH3': 'q'}
ts = data[type_col - 1].astype(str).str.strip()
keep = data[type_col - 1].notna() & ts.ne('') & ts.ne('nan')
letters = ts.map(m)
invalid = keep & letters.isna()
if invalid.any():
    raise ValueError(f"类型列存在非法取值: {sorted(set(ts[invalid]))}")
pos_upper = data[pos_col - 1].astype(str).str.upper()
letters = letters.mask(keep & pos_upper.str.contains('OCH3', na=False), 'q')
ppm = pd.to_numeric(data[ppm_col - 1], errors='coerce').round(1)

# 直接访问内部单元格字典 ws._cells，绕过逐格代理查找；只写有变化的格
cells = ws._cells
for i in np.flatnonzero(keep.to_numpy()):
    r = int(i) + 2
    tcell = cells.get((r, type_new_col))
    if tcell is None:
        tcell = ws.cell(row=r, column=type_new_col)
    if tcell.value != letters.iat[i]:
        tcell.value = letters.iat[i]
    pcell = cells.get((r, ppm_col))
    if pcell is not None and pcell.value is not None and pcell.value != ppm.iat[i]:
        pcell.value = ppm.iat[i]

# 空行按连续块自下而上批量删除（delete_rows 支持一次删多行）
empty_rows = np.flatnonzero(~keep.to_numpy()) + 2  # 工作表行号
if empty_rows.size:
    gap = np.flatnonzero(np.diff(empty_rows) > 1)
    starts = np.r_[0, gap + 1]
    ends = np.r_[gap, empty_rows.size - 1]
    for s, e in zip(starts[::-1], ends[::-1]):
        ws.delete_rows(int(empty_rows[s]), int(e - s + 1))

wb.save(SRC)
print(f'openpyxl 总耗时 {time.perf_counter() - t0:.2f} s')

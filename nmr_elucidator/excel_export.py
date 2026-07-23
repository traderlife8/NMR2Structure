"""化学家 Excel 风格 .xlsx 导出 (Part 3 Phase 6).

对标化学家论文 Table 1 的 Excel 模板:
  - 顶部 3 行元信息 (HRESIMS / [α] / IR)
  - 8 列表头 (No. / 13C NMR / Types / 1H NMR (mult., J in Hz) / COSY / HMBC / ROESY / Notes)
  - 数据行, 每个 H 一行
  - 表头加粗, 边框, 列宽合理
"""
from __future__ import annotations

import os
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .assignment_table_v2 import AssignmentTableV2


# 列定义: (列宽, 表头)
_COLUMNS = [
    (8, 'No.'),
    (10, '13C NMR'),
    (8, 'Types'),
    (28, '1H NMR (mult., J in Hz)'),
    (16, '1H-1H COSY'),
    (24, 'HMBC'),
    (16, 'ROESY'),
    (12, 'Notes'),
]


def write_assignment_table_v2_xlsx(
    table: AssignmentTableV2,
    output_path: str,
    rank: int = 0,
    score: Optional[float] = None,
    extra_meta: Optional[dict] = None,
) -> None:
    """把 v2 归属表写成 .xlsx 文件 (贴近化学家 Excel 模板).

    Args:
        table: AssignmentTableV2 实例
        output_path: 输出 .xlsx 路径
        rank: 候选排名 (0 表示手工归属)
        score: 候选打分
        extra_meta: 额外元信息 (如 'DBE', 'Solvent', 'Field' 等)
    """
    extra_meta = extra_meta or {}
    wb = Workbook()
    ws = wb.active
    ws.title = 'NMR归属'

    # 样式
    bold_font = Font(name='Calibri', size=11, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='888888'),
        right=Side(style='thin', color='888888'),
        top=Side(style='thin', color='888888'),
        bottom=Side(style='thin', color='888888'),
    )
    header_fill = PatternFill('solid', fgColor='D9E1F2')  # 浅蓝色
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    row_idx = 1
    # 顶部元信息 (化学家 Excel 第 1-3 行)
    if table.hresims:
        ws.cell(row=row_idx, column=1, value='HRESIMS').font = bold_font
        ws.cell(row=row_idx, column=2, value=table.hresims)
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=8)
        row_idx += 1
    if table.alpha:
        ws.cell(row=row_idx, column=1, value='[α]').font = bold_font
        ws.cell(row=row_idx, column=2, value=table.alpha)
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=8)
        row_idx += 1
    if table.ir:
        ws.cell(row=row_idx, column=1, value='IR').font = bold_font
        ws.cell(row=row_idx, column=2, value=table.ir)
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=8)
        row_idx += 1
    # 必备元信息
    ws.cell(row=row_idx, column=1, value='Rank').font = bold_font
    ws.cell(row=row_idx, column=2, value=rank)
    row_idx += 1
    ws.cell(row=row_idx, column=1, value='Score').font = bold_font
    ws.cell(row=row_idx, column=2, value=f'{score:.4f}' if score is not None else '')
    row_idx += 1
    ws.cell(row=row_idx, column=1, value='SMILES').font = bold_font
    ws.cell(row=row_idx, column=2, value=table.structure_smiles)
    ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=8)
    row_idx += 1
    ws.cell(row=row_idx, column=1, value='Confidence').font = bold_font
    ws.cell(row=row_idx, column=2, value=f'{table.confidence:.2%}')
    row_idx += 1
    for k, v in extra_meta.items():
        ws.cell(row=row_idx, column=1, value=str(k)).font = bold_font
        ws.cell(row=row_idx, column=2, value=str(v))
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=8)
        row_idx += 1
    if table.unassigned_obs_c:
        ws.cell(row=row_idx, column=1, value='未归属 13C (ppm)').font = bold_font
        ws.cell(row=row_idx, column=2,
                value='; '.join(f'{c:.2f}' for c in table.unassigned_obs_c))
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=8)
        row_idx += 1
    if table.unassigned_obs_h:
        ws.cell(row=row_idx, column=1, value='未归属 1H (ppm)').font = bold_font
        ws.cell(row=row_idx, column=2,
                value='; '.join(f'{h:.3f}' for h in table.unassigned_obs_h))
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=8)
        row_idx += 1

    # 表头
    header_row = row_idx
    for col_idx, (width, label) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=label)
        cell.font = bold_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    row_idx += 1

    # 数据行
    for r in table.rows:
        # No. (1α / 1β / 7α / 7β / 23a / 23b)
        c = ws.cell(row=row_idx, column=1, value=r.no_label)
        c.alignment = center
        c.border = thin_border
        # 13C NMR
        c = ws.cell(row=row_idx, column=2,
                    value=f'{r.c_ppm:.1f}' if r.c_ppm is not None else '')
        c.alignment = center
        c.border = thin_border
        # Types (CH/CH2/CH3/Cq)
        c = ws.cell(row=row_idx, column=3, value=r.multiplicity)
        c.alignment = center
        c.border = thin_border
        # 1H NMR (mult., J in Hz)
        if r.h_ppm is not None:
            if r.h_mult:
                h_cell = f'{r.h_ppm:.2f}, {r.h_mult}'
            else:
                h_cell = f'{r.h_ppm:.2f}'
        else:
            h_cell = ''
        c = ws.cell(row=row_idx, column=4, value=h_cell)
        c.alignment = left
        c.border = thin_border
        # COSY
        c = ws.cell(row=row_idx, column=5, value='/'.join(r.cosy_partners))
        c.alignment = left
        c.border = thin_border
        # HMBC
        c = ws.cell(row=row_idx, column=6, value='/'.join(r.hmbc_to_c))
        c.alignment = left
        c.border = thin_border
        # ROESY
        c = ws.cell(row=row_idx, column=7, value='/'.join(r.roesy_partners))
        c.alignment = left
        c.border = thin_border
        # Notes
        c = ws.cell(row=row_idx, column=8, value=r.notes)
        c.alignment = left
        c.border = thin_border
        row_idx += 1

    # 冻结表头
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    wb.save(output_path)